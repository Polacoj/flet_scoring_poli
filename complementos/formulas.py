"""
Calcula P.EVENTO, % ZONA y P. TOTAL para cada fila del listado de productividad.

Necesita:
  - variantes.json  (generado con excel_a_json.py a partir de variantes_detenidos.xlsx)
  - productividad_imputados.xlsx  (el listado)

Genera: productividad_imputados_calculado.xlsx (el original no se toca).
"""

import json
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

# ----------------------------------------------------------------------
# 1. CONFIGURACIÓN (lo único que normalmente vas a tocar)
# ----------------------------------------------------------------------
ARCHIVO_VARIANTES = Path("variantes.json")
ARCHIVO_LISTADO = Path("productividad_imputados.xlsx")
ARCHIVO_SALIDA = Path("productividad_imputados_calculado.xlsx")

# Nombres EXACTOS de las columnas del Excel (fila 1)
COL_TIPO = "TIPO DE EVENTO"
COL_BARRIO = "BARRIO"
COL_P_EVENTO = "P.EVENTO"
COL_PCT_ZONA = "% ZONA"
COL_P_TOTAL = "P. TOTAL"

# Tipos de evento del listado que se llaman distinto en la tabla de variantes.
# Formato:  "como aparece en el listado": "como aparece en variantes"
# (todo en minúsculas y sin tildes). Se completa cuando decidas la equivalencia.
ALIAS_TIPOS = {
    # La tabla de variantes todavía usa el nombre viejo de este tipo de evento.
    "colaboracion con imputado/s positiva": "colaboracion con intervencion directa",
    # INTERVENCION DESTACADA puntúa igual que DELITO COMPLEJO.
    "intervencion destacada": "delito complejo",
}

PCT_ZONA_CONTRAVENCION = (
    0  # CONTRAVENCION: % ZONA siempre 0, sin fórmula (P. TOTAL = P.EVENTO)
)


# ----------------------------------------------------------------------
# 2. FUNCIONES AUXILIARES
# ----------------------------------------------------------------------
def normalizar(texto) -> str:
    """'  Nuñez ' -> 'nunez'  (minúsculas, sin tildes, sin espacios sobrantes)."""
    texto = unicodedata.normalize("NFD", str(texto).strip().lower())
    return "".join(c for c in texto if unicodedata.category(c) != "Mn")


def calcular_p_total(p_evento: float, pct_zona: float) -> float:
    """
    Regla de P. TOTAL:
      - % ZONA = 0                  -> queda el valor de P.EVENTO
      - % ZONA = 0,25 / 0,50 / 0,75 -> P.EVENTO + (P.EVENTO * % ZONA)

    Ejemplo con P.EVENTO = 1,25:  zona 0 -> 1,25 | 0,25 -> 1,5625
                                  0,50 -> 1,875  | 0,75 -> 2,1875
    Nota: con zona = 0 la fórmula general da lo mismo (P.EVENTO + 0),
    así que no hace falta un caso aparte.
    """
    return p_evento + (p_evento * pct_zona)


# ----------------------------------------------------------------------
# 3. PROCESO PRINCIPAL
# ----------------------------------------------------------------------
def main() -> None:
    variantes = json.loads(ARCHIVO_VARIANTES.read_text(encoding="utf-8"))
    delitos = variantes["delitos"]
    barrios = variantes["barrios"]

    wb = load_workbook(ARCHIVO_LISTADO)
    ws = wb.active

    # Mapa "nombre de columna" -> "número de columna", leído de la fila 1.
    # Así el script no depende de que las columnas estén en cierta letra.
    columnas = {
        str(celda.value).strip(): celda.column
        for celda in ws[1]
        if celda.value is not None
    }

    sin_puntaje = {}  # tipos de evento que no están en la tabla -> cantidad
    sin_barrio = {}  # barrios que no están en la tabla -> cantidad
    procesadas = 0

    for fila in range(2, ws.max_row + 1):
        tipo_crudo = ws.cell(fila, columnas[COL_TIPO]).value
        barrio_crudo = ws.cell(fila, columnas[COL_BARRIO]).value
        if tipo_crudo is None:
            continue  # fila vacía

        tipo = normalizar(tipo_crudo)
        tipo = ALIAS_TIPOS.get(tipo, tipo)

        # --- Columna 1: P.EVENTO -------------------------------------
        p_evento = delitos.get(tipo)
        if p_evento is None:
            sin_puntaje[tipo_crudo] = sin_puntaje.get(tipo_crudo, 0) + 1

        # --- Columna 2: % ZONA ---------------------------------------
        if tipo == "contravencion":
            pct_zona = PCT_ZONA_CONTRAVENCION
        else:
            datos_barrio = barrios.get(normalizar(barrio_crudo))
            if datos_barrio is None:
                sin_barrio[barrio_crudo] = sin_barrio.get(barrio_crudo, 0) + 1
                pct_zona = None
            else:
                pct_zona = datos_barrio["base"]

        # --- Columna 3: P. TOTAL (solo si tenemos los dos datos) ------
        if p_evento is not None and pct_zona is not None:
            p_total = round(calcular_p_total(p_evento, pct_zona), 4)
        else:
            p_total = None

        ws.cell(fila, columnas[COL_P_EVENTO]).value = p_evento
        ws.cell(fila, columnas[COL_PCT_ZONA]).value = pct_zona
        ws.cell(fila, columnas[COL_P_TOTAL]).value = p_total
        procesadas += 1

    wb.save(ARCHIVO_SALIDA)

    # --- Resumen en pantalla ------------------------------------------
    print(f"Filas procesadas: {procesadas}")
    print(f"Archivo generado: {ARCHIVO_SALIDA}")
    if sin_puntaje:
        print("\nTipos de evento SIN puntaje (quedaron en blanco):")
        for nombre, cantidad in sin_puntaje.items():
            print(f"  - {nombre}: {cantidad} filas")
    if sin_barrio:
        print("\nBarrios NO encontrados (quedaron en blanco):")
        for nombre, cantidad in sin_barrio.items():
            print(f"  - {nombre!r}: {cantidad} filas")


if __name__ == "__main__":
    main()
