"""Cálculo de puntajes de los eventos con imputados.

Toma el listado generado por la pestaña «Productividad Imputados»
(``DATOS/productividad_imputados.xlsx``) y completa las columnas
``P.EVENTO``, ``% ZONA`` y ``P. TOTAL`` de cada fila.

Reglas
------
* ``P.EVENTO``: puntaje base según el tipo de evento, leído de
  ``variantes_detenidos.xlsx`` (sección ``DESCRIPCION``).
* ``% ZONA``: porcentaje del barrio (0, 25/100, 50/100 o 75/100) de la
  sección ``BARRIO`` de la misma tabla. Para ``CONTRAVENCION`` siempre es 0.
* ``P. TOTAL``: ``P.EVENTO + (P.EVENTO * % ZONA)``, redondeado a 4 decimales.

El listado original no se modifica: el resultado se guarda como
``DATOS/productividad_imputados_calculado.xlsx``.
"""

from __future__ import annotations

import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .comunes import LogFn

# ----------------------------------------------------------------------
# Configuración
# ----------------------------------------------------------------------
RAIZ = Path(__file__).resolve().parent.parent
DATOS_DIR = RAIZ / "DATOS"

#: Tabla de puntajes (copiada al proyecto desde «variantes detenidos.xlsx»).
RUTA_VARIANTES = RAIZ / "variantes_detenidos.xlsx"

ARCHIVO_LISTADO = "productividad_imputados.xlsx"
ARCHIVO_SALIDA = "productividad_imputados_calculado.xlsx"

#: Hojas del archivo de salida.
HOJA_LISTADO = "Listado"
HOJA_ACUMULADO = "Acumulado por oficial"

#: Orden de los tipos en la hoja acumulada (como el adjunto). Los tipos que
#: aparezcan y no estén en esta lista se agregan alfabéticamente al final.
ORDEN_TIPOS = [
    "DELITO",
    "ROBO/HURTO",
    "INTERVENCION DESTACADA",
    "CONTRAVENCION",
    "COLABORACION CON IMPUTADO/S POSITIVA",
]

#: Colores del encabezado y de la fila de totales.
COLOR_CABECERA = "FFD9E1F2"
COLOR_TOTALES = "FFF2F2F2"

# Nombres exactos de las columnas del listado.
COL_TIPO = "TIPO DE EVENTO"
COL_BARRIO = "BARRIO"
COL_P_EVENTO = "P.EVENTO"
COL_PCT_ZONA = "% ZONA"
COL_P_TOTAL = "P. TOTAL"

#: Tipos del listado que en la tabla de variantes se llaman distinto.
#: (claves ya normalizadas: minúsculas, sin tildes y con espacios simples)
ALIAS_TIPOS = {
    "colaboracion con imputado/s positiva": "colaboracion con intervencion directa",
    "intervencion destacada": "delito complejo",
}

#: Para CONTRAVENCION no se aplica el porcentaje del barrio.
TIPO_CONTRAVENCION = "contravencion"

#: Columnas que debe tener el listado de imputados.
COLUMNAS_REQUERIDAS = [COL_TIPO, COL_BARRIO]

#: Columnas que identifican al archivo maestro del CMU (para dar un error claro).
COLUMNAS_MAESTRO = {"CONTRAVENTORES", "DETENIDOS", "TIPO DE CIERRE", "ACLARACIONES"}


@dataclass
class Variantes:
    """Puntajes base por tipo de evento y porcentaje por barrio."""

    delitos: dict[str, float] = field(default_factory=dict)
    barrios: dict[str, float] = field(default_factory=dict)


@dataclass
class ResultadoFormulas:
    """Resumen del procesamiento, para el registro y el diálogo final."""

    ruta_salida: Path
    total: int = 0
    oficiales: int = 0
    por_tipo: dict[str, int] = field(default_factory=dict)
    puntaje_por_tipo: dict[str, float] = field(default_factory=dict)
    puntaje_total: float = 0.0
    sin_puntaje: dict[str, int] = field(default_factory=dict)
    sin_barrio: dict[str, int] = field(default_factory=dict)

    def resumen(self) -> str:
        lineas = [
            "✅ Cálculo completado exitosamente",
            "",
            f"📁 Archivo generado:\n{self.ruta_salida}",
            "",
            "📊 Eventos por tipo:",
        ]
        if self.por_tipo:
            lineas.extend(
                f"  • {tipo}: {cantidad}  ({self.puntaje_por_tipo.get(tipo, 0.0):.4g} pts)"
                for tipo, cantidad in self.por_tipo.items()
            )
        else:
            lineas.append("  • (sin eventos)")
        lineas.append(f"  • Total: {self.total} eventos")
        lineas.append(f"  • Oficiales: {self.oficiales}")
        lineas.append(f"  • Puntaje total: {self.puntaje_total:.4g}")
        lineas.append("")
        lineas.append(f"📄 Hojas: «{HOJA_LISTADO}» y «{HOJA_ACUMULADO}»")

        if self.sin_puntaje:
            lineas.append("")
            lineas.append("⚠ Tipos de evento sin puntaje en la tabla:")
            lineas.extend(
                f"  • {tipo}: {cantidad} filas" for tipo, cantidad in self.sin_puntaje.items()
            )
        if self.sin_barrio:
            lineas.append("")
            lineas.append("⚠ Barrios no encontrados en la tabla:")
            lineas.extend(
                f"  • {barrio}: {cantidad} filas" for barrio, cantidad in self.sin_barrio.items()
            )
        return "\n".join(lineas)


# ----------------------------------------------------------------------
# Utilidades
# ----------------------------------------------------------------------
def normalizar(valor) -> str:
    """'  Nuñez ' -> 'nunez' (minúsculas, sin tildes, espacios simples)."""
    if pd.isna(valor):
        return ""
    texto = unicodedata.normalize("NFKD", str(valor).strip().lower())
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return " ".join(texto.split())


def _a_numero(valor) -> float | None:
    """Convierte 25, 0.25 o la fracción '25/100' a número."""
    if pd.isna(valor):
        return None
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return float(valor)
    texto = str(valor).strip().replace(",", ".")
    if not texto:
        return None
    if "/" in texto:
        numerador, _, denominador = texto.partition("/")
        try:
            return float(numerador) / float(denominador)
        except (ValueError, ZeroDivisionError):
            return None
    try:
        return float(texto)
    except ValueError:
        return None


def _etiqueta_original(valor) -> str:
    """Texto legible para los avisos (conserva la grafía original)."""
    if pd.isna(valor) or str(valor).strip() == "":
        return "(vacío)"
    return str(valor).strip()


# ----------------------------------------------------------------------
# Tabla de variantes
# ----------------------------------------------------------------------
def cargar_variantes(ruta: Path | None = None) -> Variantes:
    """Lee la tabla de puntajes (secciones ``DESCRIPCION`` y ``BARRIO``)."""
    ruta = Path(ruta) if ruta is not None else RUTA_VARIANTES
    if not ruta.exists():
        raise RuntimeError(
            f"No se encontró la tabla de variantes:\n{ruta}\n\n"
            "Copiá «variantes detenidos.xlsx» a la carpeta del proyecto con el "
            f"nombre «{RUTA_VARIANTES.name}»."
        )

    tabla = pd.read_excel(ruta, sheet_name=0, header=None)
    if tabla.empty:
        raise RuntimeError(f"La tabla de variantes está vacía: {ruta}")

    claves = [normalizar(valor) for valor in tabla.iloc[:, 0]]
    try:
        inicio_delitos = claves.index("descripcion")
    except ValueError as exc:
        raise RuntimeError(
            "La tabla de variantes no tiene la sección 'DESCRIPCION' en la "
            f"primera columna: {ruta}"
        ) from exc
    try:
        inicio_barrios = claves.index("barrio")
    except ValueError as exc:
        raise RuntimeError(
            "La tabla de variantes no tiene la sección 'BARRIO' en la "
            f"primera columna: {ruta}"
        ) from exc

    variantes = Variantes()
    for fila in range(inicio_delitos + 1, inicio_barrios):
        nombre = claves[fila]
        valor = _a_numero(tabla.iat[fila, 1])
        if nombre and valor is not None:
            variantes.delitos[nombre] = valor

    for fila in range(inicio_barrios + 1, len(tabla)):
        nombre = claves[fila]
        valor = _a_numero(tabla.iat[fila, 1])
        if nombre and valor is not None:
            variantes.barrios[nombre] = valor

    if not variantes.delitos:
        raise RuntimeError(
            f"No se pudo leer ningún puntaje de la sección 'DESCRIPCION': {ruta}"
        )
    if not variantes.barrios:
        raise RuntimeError(f"No se pudo leer ningún barrio de la tabla: {ruta}")

    return variantes


# ----------------------------------------------------------------------
# Cálculo
# ----------------------------------------------------------------------
def calcular(
    df: pd.DataFrame, variantes: Variantes, log: LogFn
) -> tuple[pd.DataFrame, ResultadoFormulas]:
    """Completa P.EVENTO, % ZONA y P. TOTAL en el listado."""
    columnas = {str(col).strip().upper() for col in df.columns}
    faltantes = [col for col in COLUMNAS_REQUERIDAS if col not in columnas]
    if faltantes:
        raise RuntimeError(
            f"Al listado le faltan las columnas {faltantes}. "
            f"Columnas encontradas: {sorted(columnas)}"
        )

    df = df.copy()
    # Si el listado trae las columnas vacías (como lo genera Imputados), se
    # respetan; si no, se agregan al final.
    for col in (COL_P_EVENTO, COL_PCT_ZONA, COL_P_TOTAL):
        if col not in df.columns:
            df[col] = None

    tipos = df[COL_TIPO].map(normalizar)
    tipos_clave = tipos.map(lambda tipo: ALIAS_TIPOS.get(tipo, tipo))

    p_evento = tipos_clave.map(variantes.delitos)
    pct_barrio = df[COL_BARRIO].map(normalizar).map(variantes.barrios)
    # CONTRAVENCION no lleva porcentaje de zona; el resto usa el del barrio.
    pct_zona = pct_barrio.mask(tipos_clave.eq(TIPO_CONTRAVENCION), 0.0)

    p_total = (p_evento + p_evento * pct_zona).round(4)

    df[COL_P_EVENTO] = p_evento
    df[COL_PCT_ZONA] = pct_zona
    df[COL_P_TOTAL] = p_total

    # --- Avisos: tipos sin puntaje y barrios desconocidos -----------------
    filas_con_tipo = df[COL_TIPO].notna() & df[COL_TIPO].astype(str).str.strip().ne("")
    sin_puntaje = (
        df.loc[filas_con_tipo & p_evento.isna(), COL_TIPO]
        .map(_etiqueta_original)
        .value_counts()
        .to_dict()
    )

    con_barrio = (
        df[COL_BARRIO].notna()
        & df[COL_BARRIO].astype(str).str.strip().ne("")
        & tipos_clave.ne(TIPO_CONTRAVENCION)
    )
    sin_barrio = (
        df.loc[con_barrio & pct_barrio.isna(), COL_BARRIO]
        .map(_etiqueta_original)
        .value_counts()
        .to_dict()
    )

    por_tipo = {
        str(tipo): int(cantidad)
        for tipo, cantidad in df[COL_TIPO].value_counts().items()
    }
    puntaje_por_tipo = {
        str(tipo): float(grupo[COL_P_TOTAL].sum(skipna=True))
        for tipo, grupo in df.groupby(COL_TIPO, dropna=False)
    }

    resultado = ResultadoFormulas(
        ruta_salida=DATOS_DIR / ARCHIVO_SALIDA,
        total=len(df),
        por_tipo=por_tipo,
        puntaje_por_tipo=puntaje_por_tipo,
        puntaje_total=float(p_total.sum(skipna=True)),
        sin_puntaje={str(k): int(v) for k, v in sin_puntaje.items()},
        sin_barrio={str(k): int(v) for k, v in sin_barrio.items()},
    )

    log(f"✓ Filas calculadas: {resultado.total}")
    log(f"✓ Puntaje total: {resultado.puntaje_total:.4g}")
    if resultado.sin_puntaje:
        log("⚠ Tipos sin puntaje (quedan en blanco):")
        for tipo, cantidad in resultado.sin_puntaje.items():
            log(f"    • {tipo}: {cantidad}")
    if resultado.sin_barrio:
        log("⚠ Barrios no encontrados (quedan en blanco):")
        for barrio, cantidad in resultado.sin_barrio.items():
            log(f"    • {barrio}: {cantidad}")

    return df, resultado


# ----------------------------------------------------------------------
# Hoja acumulada por oficial
# ----------------------------------------------------------------------
_BORDE = Border(
    left=Side(style="thin", color="B7C3D0"),
    right=Side(style="thin", color="B7C3D0"),
    top=Side(style="thin", color="B7C3D0"),
    bottom=Side(style="thin", color="B7C3D0"),
)


def tipos_ordenados(df: pd.DataFrame) -> list[str]:
    """Tipos presentes en el listado, en el orden del adjunto."""
    presentes = sorted(
        {str(valor).strip() for valor in df[COL_TIPO].dropna() if str(valor).strip()}
    )
    preferidos = [normalizar(tipo) for tipo in ORDEN_TIPOS]

    def clave(tipo: str) -> tuple[int, object]:
        normalizado = normalizar(tipo)
        if normalizado in preferidos:
            return (0, preferidos.index(normalizado))
        return (1, normalizado)

    return sorted(presentes, key=clave)


def construir_acumulado(
    df: pd.DataFrame,
) -> tuple[list[str], list[dict], dict]:
    """Agrupa por oficial: cantidad y suma de P.TOTAL por tipo, más totales.

    Devuelve ``(tipos, filas, total_general)``. Las filas vienen ordenadas de
    mayor a menor por ``Total Suma de P. TOTAL``.
    """
    tipos = tipos_ordenados(df)
    base = df.copy()
    base["OFICIAL"] = (
        base["OFICIAL"].fillna("(sin oficial)").astype(str).str.strip()
    )
    base[COL_TIPO] = base[COL_TIPO].astype(str).str.strip()

    if tipos:
        conteo = (
            base.groupby(["OFICIAL", COL_TIPO], dropna=False)
            .size()
            .unstack(fill_value=0)
            .reindex(columns=tipos, fill_value=0)
        )
        suma = (
            base.groupby(["OFICIAL", COL_TIPO], dropna=False)[COL_P_TOTAL]
            .sum()
            .unstack(fill_value=0)
            .reindex(columns=tipos, fill_value=0)
        )
    else:
        indice = sorted(base["OFICIAL"].unique())
        conteo = pd.DataFrame(index=indice)
        suma = pd.DataFrame(index=indice)

    total_puntaje = suma.sum(axis=1) if tipos else pd.Series(0.0, index=conteo.index)
    oficiales = sorted(
        conteo.index,
        key=lambda oficial: (-float(total_puntaje.get(oficial, 0.0)), str(oficial)),
    )

    filas: list[dict] = []
    for oficial in oficiales:
        cantidades = {tipo: int(conteo.at[oficial, tipo]) for tipo in tipos}
        puntajes = {tipo: round(float(suma.at[oficial, tipo]), 4) for tipo in tipos}
        filas.append(
            {
                "oficial": str(oficial),
                "cantidades": cantidades,
                "puntajes": puntajes,
                "total_cantidad": sum(cantidades.values()),
                "total_puntaje": round(sum(puntajes.values()), 4),
            }
        )

    total = {
        "oficial": "Total general",
        "cantidades": {
            tipo: sum(fila["cantidades"][tipo] for fila in filas) for tipo in tipos
        },
        "puntajes": {
            tipo: round(sum(fila["puntajes"][tipo] for fila in filas), 4)
            for tipo in tipos
        },
        "total_cantidad": sum(fila["total_cantidad"] for fila in filas),
        "total_puntaje": round(sum(fila["total_puntaje"] for fila in filas), 4),
    }
    return tipos, filas, total


def _estilo_cabecera(celda) -> None:
    celda.font = Font(bold=True, size=10)
    celda.fill = PatternFill("solid", fgColor=COLOR_CABECERA)
    celda.border = _BORDE
    celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _escribir_acumulado(hoja, tipos: list[str], filas: list[dict], total: dict) -> None:
    """Escribe la tabla acumulada con el encabezado de dos filas del adjunto.

    Fila 1: ``oficial`` + ``cantidad``/``puntaje`` por tipo + totales.
    Fila 2: el nombre del tipo repetido en cada columna del par.
    """
    hoja.cell(row=1, column=1, value="oficial")
    hoja.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)

    for indice, tipo in enumerate(tipos):
        col_cantidad = 2 + indice * 2
        col_puntaje = col_cantidad + 1
        hoja.cell(row=1, column=col_cantidad, value="cantidad")
        hoja.cell(row=1, column=col_puntaje, value="puntaje")
        hoja.cell(row=2, column=col_cantidad, value=tipo)
        hoja.cell(row=2, column=col_puntaje, value=tipo)

    col_total_cuenta = 2 + len(tipos) * 2
    col_total_suma = col_total_cuenta + 1
    hoja.cell(row=1, column=col_total_cuenta, value="Total Cuenta de TIPO EVENTO")
    hoja.cell(row=1, column=col_total_suma, value="Total Suma de P. TOTAL")
    hoja.merge_cells(
        start_row=1, start_column=col_total_cuenta, end_row=2, end_column=col_total_cuenta
    )
    hoja.merge_cells(
        start_row=1, start_column=col_total_suma, end_row=2, end_column=col_total_suma
    )

    ultima_columna = col_total_suma
    for fila in hoja.iter_rows(min_row=1, max_row=2, min_col=1, max_col=ultima_columna):
        for celda in fila:
            if not isinstance(celda, MergedCell):
                _estilo_cabecera(celda)

    for numero, datos in enumerate(filas, start=3):
        hoja.cell(row=numero, column=1, value=datos["oficial"])
        for indice, tipo in enumerate(tipos):
            col_cantidad = 2 + indice * 2
            col_puntaje = col_cantidad + 1
            hoja.cell(row=numero, column=col_cantidad, value=datos["cantidades"][tipo])
            hoja.cell(
                row=numero, column=col_puntaje, value=datos["puntajes"][tipo]
            ).number_format = "0.0000"
        hoja.cell(row=numero, column=col_total_cuenta, value=datos["total_cantidad"])
        hoja.cell(
            row=numero, column=col_total_suma, value=datos["total_puntaje"]
        ).number_format = "0.0000"

    fila_total = 3 + len(filas)
    hoja.cell(row=fila_total, column=1, value=total["oficial"])
    for indice, tipo in enumerate(tipos):
        col_cantidad = 2 + indice * 2
        col_puntaje = col_cantidad + 1
        hoja.cell(row=fila_total, column=col_cantidad, value=total["cantidades"][tipo])
        hoja.cell(
            row=fila_total, column=col_puntaje, value=total["puntajes"][tipo]
        ).number_format = "0.0000"
    hoja.cell(row=fila_total, column=col_total_cuenta, value=total["total_cantidad"])
    hoja.cell(
        row=fila_total, column=col_total_suma, value=total["total_puntaje"]
    ).number_format = "0.0000"

    for fila in hoja.iter_rows(
        min_row=3, max_row=fila_total, min_col=1, max_col=ultima_columna
    ):
        for celda in fila:
            celda.border = _BORDE
            celda.alignment = Alignment(
                horizontal="left" if celda.column == 1 else "center",
                vertical="center",
            )
    for celda in hoja[fila_total]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill("solid", fgColor=COLOR_TOTALES)

    hoja.column_dimensions["A"].width = 34
    for indice in range(len(tipos) * 2):
        letra = get_column_letter(2 + indice)
        hoja.column_dimensions[letra].width = 10 if indice % 2 == 0 else 12
    hoja.column_dimensions[get_column_letter(col_total_cuenta)].width = 20
    hoja.column_dimensions[get_column_letter(col_total_suma)].width = 20
    hoja.row_dimensions[1].height = 30
    hoja.row_dimensions[2].height = 24
    hoja.freeze_panes = "B3"


def _validar_listado(df: pd.DataFrame, ruta: Path) -> None:
    columnas = {str(col).strip().upper() for col in df.columns}
    if COLUMNAS_MAESTRO & columnas:
        raise RuntimeError(
            f"«{ruta.name}» es el archivo maestro del CMU, no el listado.\n\n"
            "Ejecutá primero la pestaña «Productividad Imputados» y elegí "
            f"«{ARCHIVO_LISTADO}»."
        )
    faltantes = [col for col in COLUMNAS_REQUERIDAS if col not in columnas]
    if faltantes:
        raise RuntimeError(
            f"«{ruta.name}» no parece el listado de imputados: faltan las "
            f"columnas {faltantes}.\n\n"
            "Ejecutá primero la pestaña «Productividad Imputados» y elegí "
            f"«{ARCHIVO_LISTADO}»."
        )


def procesar(ruta_listado: Path | str, log: LogFn) -> ResultadoFormulas:
    """Calcula los puntajes del listado y guarda ``..._calculado.xlsx``."""
    ruta_listado = Path(ruta_listado)
    log("═══════════════════════════════════════════════════════════════════════")
    log("🧮 INICIANDO CÁLCULO DE PUNTAJES (IMPUTADOS)")
    log("═══════════════════════════════════════════════════════════════════════")
    log(f"📂 Listado: {ruta_listado.name}")

    log("→ Cargando listado...")
    df = pd.read_excel(ruta_listado, engine="openpyxl")
    _validar_listado(df, ruta_listado)
    log(f"✓ Listado cargado — {df.shape[0]} filas × {df.shape[1]} columnas")

    log(f"→ Cargando tabla de variantes: {RUTA_VARIANTES.name}...")
    variantes = cargar_variantes()
    log(
        f"✓ Variantes cargadas — {len(variantes.delitos)} tipos de evento, "
        f"{len(variantes.barrios)} barrios"
    )

    log("→ Calculando P.EVENTO, % ZONA y P. TOTAL...")
    df_calculado, resultado = calcular(df, variantes, log)

    log("→ Armando la hoja acumulada por oficial...")
    tipos, filas, total_oficiales = construir_acumulado(df_calculado)
    resultado.oficiales = len(filas)

    DATOS_DIR.mkdir(exist_ok=True)
    salida = DATOS_DIR / ARCHIVO_SALIDA
    with pd.ExcelWriter(salida, engine="openpyxl") as escritor:
        df_calculado.to_excel(escritor, sheet_name=HOJA_LISTADO, index=False)
        hoja = escritor.book.create_sheet(HOJA_ACUMULADO)
        _escribir_acumulado(hoja, tipos, filas, total_oficiales)
    resultado.ruta_salida = salida
    log(
        f"✓ Hoja «{HOJA_ACUMULADO}»: {len(filas)} oficiales, "
        f"{len(tipos)} tipos de evento"
    )

    log("───────────────────────────────────────────────────────────────────────")
    for tipo, cantidad in resultado.por_tipo.items():
        puntos = resultado.puntaje_por_tipo.get(tipo, 0.0)
        log(f"  • {tipo}: {cantidad} eventos → {puntos:.4g} pts")
    log(f"✓ Puntaje total: {resultado.puntaje_total:.4g}")
    log(f"✓ Hojas: «{HOJA_LISTADO}» y «{HOJA_ACUMULADO}»")
    log(f"✓ Archivo guardado en: {salida}")
    log("═══════════════════════════════════════════════════════════════════════")

    return resultado
