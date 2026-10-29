"""Reconciliación nombres → ID (LP) del Diccionario Maestro de OFICIALES.

Portado de ``diccionario_OFICIALES.pyw`` (PyQt6). El flujo es el mismo:

1. Se construye un mapa ``APELLIDO Y NOMBRE`` → ``LP`` con la hoja Sheet1.
2. Se normaliza la columna ``OFICIAL`` del archivo a analizar y se reemplaza
   cada nombre por su ID.
3. Si el porcentaje de coincidencia no es 100 %, se agregan al diccionario los
   nombres que faltaban (con backup ``ANTIGUO_DICCIONARIO.xlsx``) y se aborta
   para que el usuario complete los IDs.
4. El proceso inverso lee la hoja ``validacion`` y reemplaza IDs por nombres.
"""

from __future__ import annotations

import shutil
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

from .comunes import (
    LogFn,
    buscar_hoja,
    leer_hoja,
    listar_hojas,
    normalizar_columnas,
    normalizar_nombre,
    ruta_procesado,
    ruta_revertido,
    serie_id,
)

# --- Encabezados de la hoja del diccionario (Sheet1) ---
COL_ID_DICCIONARIO = "LP"
COL_NOMBRE_DICCIONARIO = "APELLIDO Y NOMBRE"
COLUMNAS_DICCIONARIO = [COL_NOMBRE_DICCIONARIO, COL_ID_DICCIONARIO]

# --- Encabezados requeridos en el archivo a analizar ---
COLUMNAS_ANALISIS = ["OFICIAL"]

# --- Hojas auxiliares ---
HOJA_PREFERIDA_DICCIONARIO = "SHEET1"
HOJA_VALIDACION = "VALIDACION"

# La hoja VALIDACION de los diccionarios reales usa otros encabezados que la
# hoja Sheet1. Se aceptan ambos juegos de nombres: antes el proceso inverso
# buscaba "LP" y "APELLIDO Y NOMBRE" dentro de VALIDACION (que tiene "LP/DNI"
# y "NOMBRE Y APELLIDO"), por lo que fallaba siempre.
COLS_ID_VALIDACION = ("LP/DNI", "LP", "DNI")
COLS_NOMBRE_VALIDACION = ("NOMBRE Y APELLIDO", "APELLIDO Y NOMBRE")

BACKUP_DICCIONARIO = "ANTIGUO_DICCIONARIO.xlsx"


class ProcesoIncompleto(RuntimeError):
    """Coincidencia menor al 100 %: se actualizó el diccionario o faltan IDs."""


@dataclass
class ResultadoReconciliacion:
    """Estadísticas de la reconciliación de la columna ``OFICIAL``."""

    celdas_evaluadas: int = 0
    coincidencias: int = 0
    no_encontrados: int = 0
    vacias: int = 0
    sin_id: int = 0
    nombres_nuevos: dict[str, list[int]] = field(default_factory=dict)
    nombres_sin_id: dict[str, list[int]] = field(default_factory=dict)

    @property
    def porcentaje(self) -> float:
        if self.celdas_evaluadas == 0:
            return 100.0
        return self.coincidencias / self.celdas_evaluadas * 100.0


def construir_mapa_nombres(df_diccionario: pd.DataFrame) -> dict[str, str]:
    """``{APELLIDO Y NOMBRE normalizado: LP}`` a partir de la hoja Sheet1."""
    df = normalizar_columnas(df_diccionario)
    faltantes = [c for c in (COL_NOMBRE_DICCIONARIO, COL_ID_DICCIONARIO) if c not in df.columns]
    if faltantes:
        raise RuntimeError(
            f"El Diccionario Maestro debe contener las columnas "
            f"'{COL_ID_DICCIONARIO}' y '{COL_NOMBRE_DICCIONARIO}'. "
            f"Faltan: {faltantes}. Columnas encontradas: {list(df.columns)}"
        )
    nombres = df[COL_NOMBRE_DICCIONARIO].map(normalizar_nombre)
    ids = serie_id(df[COL_ID_DICCIONARIO])
    mask = nombres != ""
    return dict(zip(nombres[mask], ids[mask], strict=False))


def _columna_presente(df: pd.DataFrame, candidatas: tuple[str, ...]) -> str | None:
    presentes = set(df.columns)
    return next((c for c in candidatas if c in presentes), None)


def construir_mapa_inverso(df_validacion: pd.DataFrame) -> dict[str, str]:
    """``{LP/DNI: NOMBRE Y APELLIDO}`` a partir de la hoja VALIDACION."""
    df = normalizar_columnas(df_validacion)
    col_id = _columna_presente(df, COLS_ID_VALIDACION)
    col_nombre = _columna_presente(df, COLS_NOMBRE_VALIDACION)
    if col_id is None or col_nombre is None:
        raise RuntimeError(
            f"La hoja 'VALIDACION' no contiene las columnas esperadas. "
            f"Se buscó un ID entre {list(COLS_ID_VALIDACION)} y un nombre entre "
            f"{list(COLS_NOMBRE_VALIDACION)}. Columnas encontradas: {list(df.columns)}"
        )
    ids = serie_id(df[col_id])
    nombres = df[col_nombre].fillna("").astype(str).str.strip()
    mask = ids != ""
    return dict(zip(ids[mask], nombres[mask], strict=False))


def leer_validacion(ruta_diccionario: Path, log: LogFn) -> pd.DataFrame:
    """Lee la hoja VALIDACION del diccionario (ignorando mayúsculas)."""
    nombres_hojas = listar_hojas(ruta_diccionario)
    hoja = buscar_hoja(nombres_hojas, HOJA_VALIDACION)
    if hoja is None:
        raise RuntimeError(
            f"No se encontró la hoja '{HOJA_VALIDACION}' en "
            f"{ruta_diccionario.name}.\nHojas disponibles: {nombres_hojas}"
        )
    df = normalizar_columnas(leer_hoja(ruta_diccionario, hoja))
    log(f"  Hoja encontrada: '{hoja}' — {len(df)} filas.")
    return df


def guardar_diccionario_con_bak(df_actualizado: pd.DataFrame, ruta_original: Path) -> Path:
    """Copia el original a ``ANTIGUO_DICCIONARIO.xlsx`` y reescribe Sheet1.

    Se conservan las demás hojas (por ejemplo ``validacion``).
    """
    ruta_bak = ruta_original.parent / BACKUP_DICCIONARIO
    shutil.copy2(ruta_original, ruta_bak)

    wb = load_workbook(ruta_original)
    nombre_sheet1 = buscar_hoja(wb.sheetnames, HOJA_PREFERIDA_DICCIONARIO)
    ws = wb[nombre_sheet1] if nombre_sheet1 else wb.create_sheet("Sheet1")
    if ws.max_row:  # en una hoja vacía delete_rows(1, 0) da problemas
        ws.delete_rows(1, ws.max_row)

    df_out = df_actualizado.where(df_actualizado.notna(), "")
    for fila in dataframe_to_rows(df_out, index=False, header=True):
        ws.append(fila)

    wb.save(ruta_original)
    return ruta_bak


def _agregar_nombres(df_diccionario: pd.DataFrame, nombres: list[str]) -> pd.DataFrame:
    """Agrega una fila vacía por cada nombre nuevo (sin ID, para completar)."""
    if not nombres:
        return df_diccionario
    nuevas = [{col: "" for col in df_diccionario.columns} for _ in nombres]
    for fila, nombre in zip(nuevas, nombres, strict=True):
        fila[COL_NOMBRE_DICCIONARIO] = nombre
    return pd.concat([df_diccionario, pd.DataFrame(nuevas)], ignore_index=True)


def _informar_incompleto(
    stats: ResultadoReconciliacion,
    ruta_diccionario: Path,
    log: LogFn,
) -> None:
    if stats.nombres_sin_id:
        log("Hay nombres del archivo que están en el diccionario pero sin LP/DNI:")
        for nombre, filas in stats.nombres_sin_id.items():
            log(f"  Sin ID: '{nombre}' en filas {filas}")
    if stats.nombres_nuevos:
        log("Coincidencia incompleta. Actualizando diccionario original con los faltantes...")
        for nombre, filas in stats.nombres_nuevos.items():
            log(f"  Nombre nuevo: '{nombre}' en filas {filas}")


def reconciliar(
    df_trabajo: pd.DataFrame,
    df_diccionario: pd.DataFrame,
    ruta_trabajo: Path,
    ruta_diccionario: Path,
    log: LogFn,
) -> tuple[Path, ResultadoReconciliacion]:
    """Reemplaza los nombres de la columna ``OFICIAL`` por su LP.

    Devuelve ``(ruta_del_PROCESADO, estadísticas)``. Si la coincidencia no es
    del 100 %, actualiza el diccionario (con backup) y lanza
    :class:`ProcesoIncompleto`.
    """
    df_trabajo = normalizar_columnas(df_trabajo)
    df_diccionario = normalizar_columnas(df_diccionario)

    if COLUMNAS_ANALISIS[0] not in df_trabajo.columns:
        raise RuntimeError(
            f"La hoja del archivo a analizar no contiene la columna "
            f"'{COLUMNAS_ANALISIS[0]}'."
        )

    log(f"Iniciando reconciliación de {ruta_trabajo.name}...")
    mapa = construir_mapa_nombres(df_diccionario)
    df_mod = df_trabajo.copy().reset_index(drop=True)

    stats = ResultadoReconciliacion()
    normalizados = df_mod[COLUMNAS_ANALISIS[0]].map(normalizar_nombre)
    vacias = normalizados.eq("")
    ids = normalizados.map(mapa)
    en_mapa = ~vacias & ids.notna()
    coincidencias = en_mapa & ids.ne("")
    sin_id = en_mapa & ids.eq("")
    no_encontrados = ~vacias & ids.isna()

    stats.vacias = int(vacias.sum())
    stats.coincidencias = int(coincidencias.sum())
    stats.no_encontrados = int(no_encontrados.sum())
    stats.sin_id = int(sin_id.sum())
    # Las celdas vacías no cuentan para el porcentaje (igual que antes).
    stats.celdas_evaluadas = int((~vacias).sum())

    for nombre in normalizados[no_encontrados].unique():
        filas = [int(i) + 2 for i in normalizados.index[normalizados.eq(nombre) & no_encontrados]]
        stats.nombres_nuevos[str(nombre)] = filas
    for nombre in normalizados[sin_id].unique():
        filas = [int(i) + 2 for i in normalizados.index[normalizados.eq(nombre) & sin_id]]
        stats.nombres_sin_id[str(nombre)] = filas

    log(
        f"{ruta_trabajo.name}: {stats.celdas_evaluadas} celdas evaluadas, "
        f"{stats.coincidencias} coincidencias, {stats.no_encontrados} no encontradas, "
        f"{stats.sin_id} sin ID, {stats.vacias} vacías, {stats.porcentaje:.1f} %."
    )

    if stats.porcentaje < 100.0:
        _informar_incompleto(stats, ruta_diccionario, log)
        if stats.nombres_nuevos:
            df_act = _agregar_nombres(df_diccionario, list(stats.nombres_nuevos))
            ruta_bak = guardar_diccionario_con_bak(df_act, ruta_diccionario)
            log(f"  ✔ Backup guardado: {ruta_bak.name}")
            log(f"  ✔ Diccionario original actualizado: {ruta_diccionario.name}")

        partes = [f"Coincidencia incompleta ({stats.porcentaje:.1f} %)."]
        if stats.nombres_nuevos:
            partes.append(
                f"Se actualizó '{ruta_diccionario.name}' con los nombres faltantes."
            )
        if stats.nombres_sin_id:
            partes.append(
                "Hay entradas del diccionario sin LP (no se puede reemplazar el nombre)."
            )
        partes.append("Completá la columna ID y volvé a correr.")
        raise ProcesoIncompleto(" ".join(partes))

    df_mod.loc[coincidencias, COLUMNAS_ANALISIS[0]] = ids[coincidencias].astype(str)
    salida = ruta_procesado(ruta_trabajo)
    df_mod.to_excel(salida, index=False)
    log(f"Archivo analizado y procesado guardado en: {salida.name}")
    return salida, stats


def revertir(
    ruta_procesada: Path,
    columnas: list[str],
    mapa_inverso: dict[str, str],
    etiqueta: str,
    log: LogFn,
) -> Path:
    """Reemplaza los IDs de un ``_PROCESADO`` por nombres y guarda ``_REVERTIDO``."""
    log(f"\nProcesando {etiqueta}: {ruta_procesada.name}...")
    hojas = listar_hojas(ruta_procesada)
    if not hojas:
        raise RuntimeError(f"{etiqueta} no tiene hojas válidas.")
    df = normalizar_columnas(leer_hoja(ruta_procesada, hojas[0])).reset_index(drop=True)

    reemplazados = 0
    vacias = 0
    ids_sin_match: list[str] = []

    for col in columnas:
        if col not in df.columns:
            log(f"  ⚠ Columna '{col}' no encontrada en {ruta_procesada.name}, se omite.")
            continue
        valores = df[col].fillna("").astype(str).str.strip()
        vacio = valores.eq("") | valores.str.lower().eq("nan")
        match = ~vacio & valores.isin(mapa_inverso)
        sin_match = ~vacio & ~valores.isin(mapa_inverso)
        df.loc[match, col] = valores[match].map(mapa_inverso)
        reemplazados += int(match.sum())
        vacias += int(vacio.sum())
        for idx, valor in valores[sin_match].items():
            # +2 = una fila de encabezado y el índice base 0
            ids_sin_match.append(f"{col}:fila {int(idx) + 2} → '{valor}'")

    log(f"  ✔ Reemplazados:  {reemplazados}")
    log(f"  ○ Vacíos:        {vacias}")
    if ids_sin_match:
        log(f"  ⚠ Sin coincidencia en VALIDACION: {len(ids_sin_match)}")
        for detalle in ids_sin_match:
            log(f"     {detalle}")

    salida = ruta_revertido(ruta_procesada)
    df.to_excel(salida, index=False)
    log(f"  ✔ Guardado en: {salida.name}")
    return salida
